"""Tests for the Excel datasource, including the openpyxl accessors used on Python 3.12."""
import types

import openpyxl

from panorama_elt.xls_datasource.xls_datasource import XLSDatasource


class FakeSheet:
    """Backs ``sheet.cell(row, column).value`` with a sparse grid dict."""

    def __init__(self, grid):
        self.grid = grid

    def cell(self, row, column):
        return types.SimpleNamespace(value=self.grid.get((row, column)))


class FakeWorkbook:
    def __init__(self, sheets):
        self._sheets = {name: FakeSheet(grid) for name, grid in sheets.items()}
        self.sheetnames = list(sheets.keys())
        self.closed = False

    def __getitem__(self, name):
        return self._sheets[name]

    def close(self):
        self.closed = True


class FakeDatalake:
    def __init__(self):
        self.uploads = []

    def upload_table_from_file(self, **kwargs):
        self.uploads.append(kwargs)


def patch_workbook(monkeypatch, sheets):
    """Patch openpyxl.load_workbook to hand back fresh fake workbooks, recording each."""
    instances = []

    def load_workbook(_location):
        workbook = FakeWorkbook(sheets)
        instances.append(workbook)
        return workbook

    monkeypatch.setattr(openpyxl, "load_workbook", load_workbook)
    return instances


def test_get_tables_returns_sheet_names_and_closes(monkeypatch):
    instances = patch_workbook(monkeypatch, {"Sheet1": {}, "Sheet2": {}})
    ds = XLSDatasource(datalake=FakeDatalake(), datasource_settings={"location": "/data.xlsx"})

    assert ds.get_tables() == ["Sheet1", "Sheet2"]
    assert instances[0].closed


def test_get_fields_reads_header_via_cell_accessors(monkeypatch):
    grid = {(1, 1): "username", (1, 2): "email", (1, 3): None}
    instances = patch_workbook(monkeypatch, {"Sheet1": grid})
    ds = XLSDatasource(datalake=FakeDatalake(), datasource_settings={"location": "/data.xlsx"})

    assert ds.get_fields("Sheet1") == [
        {"name": "username", "type": "string"},
        {"name": "email", "type": "string"},
    ]
    assert instances[0].closed


def test_get_fields_returns_configured_definition(monkeypatch):
    # configured fields must be returned without opening the workbook
    monkeypatch.setattr(openpyxl, "load_workbook", lambda _location: (_ for _ in ()).throw(
        AssertionError("workbook should not be opened")))
    ds = XLSDatasource(
        datalake=FakeDatalake(),
        datasource_settings={
            "location": "/data.xlsx",
            "tables": [{"name": "Sheet1", "fields": [{"name": "id"}, {"name": "name"}]}],
        },
    )
    assert ds.get_fields("Sheet1") == ["id", "name"]


def test_test_connections_reports_presence(tmp_path):
    present = tmp_path / "book.xlsx"
    present.write_text("x")
    ds_ok = XLSDatasource(datalake=FakeDatalake(), datasource_settings={"location": str(present)})
    assert ds_ok.test_connections() == {"XLS": "OK"}

    ds_missing = XLSDatasource(datalake=FakeDatalake(), datasource_settings={"location": str(tmp_path / "no.xlsx")})
    assert ds_missing.test_connections()["XLS"].startswith("File ")


def test_extract_and_load_writes_csv_and_uploads(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    grid = {
        (1, 1): "a", (1, 2): "b", (1, 3): None,   # header
        (2, 1): 1, (2, 2): 2,                       # data row
        (3, 1): None, (3, 2): None,                 # all-None row stops reading
    }
    patch_workbook(monkeypatch, {"Sheet1": grid})
    datalake = FakeDatalake()
    ds = XLSDatasource(datalake=datalake, datasource_settings={"location": "/data.xlsx"})

    ds.extract_and_load()

    assert datalake.uploads == [{
        "filename": "Sheet1.csv",
        "table": "Sheet1",
        "update_partitions": True,
    }]
    # temporary csv is cleaned up after upload
    assert not (tmp_path / "Sheet1.csv").exists()


def test_extract_and_load_skips_unselected_tables(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    patch_workbook(monkeypatch, {"Sheet1": {(1, 1): "a", (1, 2): None}})
    datalake = FakeDatalake()
    ds = XLSDatasource(datalake=datalake, datasource_settings={"location": "/data.xlsx"})

    ds.extract_and_load(selected_tables="OtherSheet")
    assert datalake.uploads == []
